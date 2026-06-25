import csv
import io
import json
import sqlite3
import uuid
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import Response
from openpyxl import Workbook
from sse_starlette.sse import EventSourceResponse

from ..auth.deps import get_current_user, require_role
from ..auth.service import verify_project_access
from ..copilot.biomaterials_skill import (
    BIOMATERIALS_SYSTEM_PROMPT,
    DOMAIN_REFUSAL,
    PROGRAMMABLE_BIOMATERIALS_SKILL,
    is_programmable_biomaterials_question,
)
from ..db import get_connection
from ..repositories import (
    automation,
    catalog,
    experiment_plans,
    knowledge,
    literature,
    literature_subscriptions,
    research_execution,
    research_planner,
)
from ..schemas import (
    CandidateExplanationRequest,
    ClaimRelationDetectRequest,
    ClaimReviewRequest,
    ClusterDraftRequest,
    CopilotChatRequest,
    CopilotConfigUpdateRequest,
    LiteratureIngestRequest,
    LiteratureSubscriptionRequest,
    MarkdownResearchSourceRequest,
    EvidenceReviewRequest,
    ExperimentPlanUpdateRequest,
    ExperimentStepUpdateRequest,
    ResearchFindingReviewRequest,
    ResultInterpretationRequest,
    ResearchBriefCreateRequest,
    ResearchPlanRequest,
    RoutePlanRequest,
    SequenceComparisonRequest,
    StructureComparisonRequest,
    WorkflowPlanMaterializeRequest,
)
from ..settings import get_settings
from ..utils.response import envelope
from ..copilot import cluster

router = APIRouter(prefix="/copilot")

SKILL_KEYWORDS = {
    "programmable-biomaterials-expert": PROGRAMMABLE_BIOMATERIALS_SKILL["trigger"],
    "workflow-adjust": ["workflow", "route", "threshold", "工作流"],
    "result-interpret": ["bli", "sec", "experiment", "实验", "lab"],
    "query-candidates": ["candidate", "rank", "anchor", "候选", "c4361"],
    "structure-explain": ["structure", "pdb", "interface", "结构"],
    "paper-reader": ["paper", "literature", "citation", "论文"],
}


def match_skill(text: str) -> str:
    lower = text.lower()
    for skill, tokens in SKILL_KEYWORDS.items():
        if any(token.lower() in lower for token in tokens):
            return skill
    return "general"


def top_candidates_message(connection: sqlite3.Connection, project_id: str | None) -> str:
    if not project_id:
        return "Select a project to query ranked candidates."
    items, _ = catalog.list_project_candidates_filtered(
        connection,
        project_id,
        decision="Anchor,Order,Retest",
        limit=5,
    )
    if not items:
        return "No prioritized candidates found for this project."
    names = ", ".join(item["candidate_id"] for item in items[:3])
    return (
        f"Prioritize {names}. They balance interface score, pLDDT, and lower aggregation risk "
        "based on the current demo dataset."
    )


def workflow_adjust_message(connection: sqlite3.Connection, project_id: str | None) -> str:
    if not project_id:
        return "Raise solubility threshold to 88 and cap each scaffold family at 6 ordered variants."
    package = catalog.get_project_delivery_package(connection, project_id)
    constraints = (package or {}).get("redesign_constraints") or {}
    if isinstance(constraints, str):
        constraints = json.loads(constraints)
    preserve = constraints.get("preserve_candidate", "PD1Binder_c4361")
    return (
        f"Preserve {preserve} motif, raise solubility threshold to 88, add hydrophobic patch penalty, "
        "and cap each scaffold family at 6 ordered variants."
    )


def result_interpret_message(connection: sqlite3.Connection, project_id: str | None) -> str:
    if not project_id:
        return "9/48 BLI-positive; SEC aggregation is the main QC loss in the PD-1 demo."
    summary = catalog.get_project_results_summary(connection, project_id)
    return (
        f"{summary['hit_rate_label']}; best BLI Kd is {summary['best_kd']}. "
        f"{summary['main_failure']} explains most QC loss. {summary['decision_detail']}"
    )


def knowledge_message(connection: sqlite3.Connection, query: str) -> str:
    items = knowledge.search_entries(connection, query, limit=3)
    if not items:
        return (
            "I do not have a matching curated knowledge entry yet. Add a knowledge entry for this method, "
            "model, algorithm, protein, assay, or biomaterial property before relying on Copilot."
        )
    lines = []
    for item in items:
        lines.append(f"- {item['title']}: {item['summary']}")
    return "Relevant programmable biomaterials knowledge:\n" + "\n".join(lines)


def literature_message(connection: sqlite3.Connection, query: str) -> str:
    items = literature.search_library(connection, query, limit=5)
    if not items:
        return (
            "The local literature library has no matching indexed evidence yet. "
            "An administrator can ingest a Europe PMC query first."
        )
    lines = []
    seen: set[tuple[str, str | None]] = set()
    for item in items:
        identity = (item["document_id"], item.get("claim_id"))
        if identity in seen:
            continue
        seen.add(identity)
        citation = item.get("doi") or (
            f"PMID:{item['pmid']}" if item.get("pmid") else item.get("external_id")
        )
        statement = item.get("statement") or item.get("abstract_text") or "Metadata record"
        lines.append(
            f"- {item['title']} ({item.get('publication_year') or 'n.d.'}; {citation}): "
            f"{statement[:320]}"
        )
    return "Indexed literature evidence:\n" + "\n".join(lines)


def _rule_based_chat(connection: sqlite3.Connection, payload: CopilotChatRequest) -> dict:
    last_message = payload.messages[-1].content if payload.messages else ""
    skill = payload.skill or match_skill(last_message)

    if skill == "programmable-biomaterials-expert":
        message = knowledge_message(connection, last_message)
    elif skill == "workflow-adjust":
        message = workflow_adjust_message(connection, payload.project_id)
    elif skill == "result-interpret":
        message = result_interpret_message(connection, payload.project_id)
    elif skill == "query-candidates":
        message = top_candidates_message(connection, payload.project_id)
    elif skill == "paper-reader":
        message = literature_message(connection, last_message)
    elif skill == "structure-explain":
        message = (
            "Explain interface contacts, hydrophobic patches, and developability risks using uploaded "
            "PDB context and candidate structure files."
        )
    else:
        message = top_candidates_message(connection, payload.project_id)

    return {
        "mode": "rule_based_demo",
        "message": message,
        "skill_used": skill,
        "structured": {
            "echo": last_message,
            "project_id": payload.project_id,
        },
    }


def _domain_refusal(project_id: str | None, last_message: str) -> dict:
    return {
        "mode": "domain_guard",
        "message": DOMAIN_REFUSAL,
        "skill_used": "programmable-biomaterials-expert",
        "structured": {
            "echo": last_message,
            "project_id": project_id,
            "allowed_scope": PROGRAMMABLE_BIOMATERIALS_SKILL["description"],
        },
    }


def _resolve_chat(connection: sqlite3.Connection, payload: CopilotChatRequest) -> dict:
    last_message = payload.messages[-1].content if payload.messages else ""
    if last_message and not is_programmable_biomaterials_question(last_message):
        return _domain_refusal(payload.project_id, last_message)
    settings = get_settings()
    if settings.llm_api_key:
        from ..copilot.service import chat_with_llm

        return chat_with_llm(connection, payload)
    return _rule_based_chat(connection, payload)


def _ensure_project_access(connection: sqlite3.Connection, user: dict, project_id: str | None) -> None:
    if project_id and not verify_project_access(connection, user, project_id):
        raise HTTPException(status_code=403, detail="forbidden")


def _copilot_config_payload() -> dict:
    settings = get_settings()
    return {
        "llm_api_base": settings.llm_api_base,
        "llm_model": settings.llm_model,
        "api_key_configured": bool(settings.llm_api_key),
        "api_key_preview": f"...{settings.llm_api_key[-4:]}" if settings.llm_api_key else None,
        "system_scope": "programmable_biomaterials_only",
        "system_prompt": BIOMATERIALS_SYSTEM_PROMPT,
    }


@router.get("/config")
def get_copilot_config(_admin: dict = Depends(require_role("admin"))):
    return envelope(_copilot_config_payload())


@router.put("/config")
def update_copilot_config(
    payload: CopilotConfigUpdateRequest,
    _admin: dict = Depends(require_role("admin")),
):
    settings = get_settings()
    if payload.llm_api_base is not None:
        settings.llm_api_base = payload.llm_api_base.strip()
    if payload.llm_model is not None:
        settings.llm_model = payload.llm_model.strip()
    if payload.llm_api_key is not None:
        settings.llm_api_key = payload.llm_api_key.strip()
    return envelope(_copilot_config_payload())


@router.post("/config/test")
def test_copilot_config(
    _admin: dict = Depends(require_role("admin")),
):
    settings = get_settings()
    if not settings.llm_api_key:
        return envelope({
            "connected": False,
            "model": settings.llm_model,
            "reason": "no_api_key",
        })
    try:
        from ..copilot.provider import get_llm_provider

        response = get_llm_provider().chat([
            {"role": "user", "content": "Reply with exactly BDA_OK."}
        ])
        return envelope({
            "connected": True,
            "model": settings.llm_model,
            "sample": response.content[:120],
        })
    except Exception as exc:
        return envelope({
            "connected": False,
            "model": settings.llm_model,
            "reason": str(exc)[:500],
        })


@router.get("/knowledge")
def list_knowledge_entries(
    q: str | None = Query(default=None, max_length=200),
    category: str | None = Query(default=None, max_length=80),
    limit: int = Query(default=20, ge=1, le=100),
    offset: int = Query(default=0, ge=0),
    connection: sqlite3.Connection = Depends(get_connection),
    _user: dict = Depends(get_current_user),
):
    if q:
        items = knowledge.search_entries(connection, q, category=category, limit=limit)
        return envelope({"items": items, "total": len(items), "limit": limit, "offset": 0, "query": q})
    items, total = knowledge.list_entries(connection, category=category, limit=limit, offset=offset)
    return envelope({"items": items, "total": total, "limit": limit, "offset": offset, "query": q})


@router.get("/knowledge/{knowledge_entry_id}")
def get_knowledge_entry(
    knowledge_entry_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    _user: dict = Depends(get_current_user),
):
    item = knowledge.get_entry(connection, knowledge_entry_id)
    if item is None:
        raise HTTPException(status_code=404, detail="knowledge_entry_not_found")
    return envelope(item)


@router.post("/literature/ingest")
def ingest_literature(
    payload: LiteratureIngestRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    _admin: dict = Depends(require_role("admin")),
):
    from ..services.literature_ingestion import ingest_europe_pmc_query

    try:
        result = ingest_europe_pmc_query(
            connection,
            payload.query,
            limit=payload.limit,
            fetch_full_text=payload.fetch_full_text,
            extract_claims=payload.extract_claims,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail=f"literature_ingestion_failed:{str(exc)[:300]}",
        ) from exc
    return envelope(result)


@router.get("/literature/subscriptions")
def list_literature_subscriptions(
    connection: sqlite3.Connection = Depends(get_connection),
    _admin: dict = Depends(require_role("admin")),
):
    items = literature_subscriptions.list_subscriptions(connection)
    return envelope({"items": items, "total": len(items)})


@router.post("/literature/subscriptions")
def create_literature_subscription(
    payload: LiteratureSubscriptionRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    admin: dict = Depends(require_role("admin")),
):
    subscription_id = f"subscription_{uuid.uuid4().hex[:12]}"
    connection.execute(
        """
        INSERT INTO literature_subscriptions (
            subscription_id, name, query, enabled, interval_hours,
            result_limit, fetch_full_text, extract_claims, created_by
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            subscription_id,
            payload.name,
            payload.query,
            int(payload.enabled),
            payload.interval_hours,
            payload.result_limit,
            int(payload.fetch_full_text),
            int(payload.extract_claims),
            admin["user_id"],
        ),
    )
    return envelope(literature_subscriptions.get_subscription(connection, subscription_id))


@router.patch("/literature/subscriptions/{subscription_id}")
def update_literature_subscription(
    subscription_id: str,
    payload: LiteratureSubscriptionRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    _admin: dict = Depends(require_role("admin")),
):
    if literature_subscriptions.get_subscription(connection, subscription_id) is None:
        raise HTTPException(status_code=404, detail="literature_subscription_not_found")
    connection.execute(
        """
        UPDATE literature_subscriptions
        SET name=?, query=?, enabled=?, interval_hours=?, result_limit=?,
            fetch_full_text=?, extract_claims=?, updated_at=CURRENT_TIMESTAMP
        WHERE subscription_id=?
        """,
        (
            payload.name,
            payload.query,
            int(payload.enabled),
            payload.interval_hours,
            payload.result_limit,
            int(payload.fetch_full_text),
            int(payload.extract_claims),
            subscription_id,
        ),
    )
    return envelope(literature_subscriptions.get_subscription(connection, subscription_id))


@router.post("/literature/subscriptions/{subscription_id}/run")
def run_literature_subscription(
    subscription_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    _admin: dict = Depends(require_role("admin")),
):
    item = literature_subscriptions.get_subscription(connection, subscription_id)
    if item is None:
        raise HTTPException(status_code=404, detail="literature_subscription_not_found")
    from ..services.literature_subscription_service import run_subscription

    return envelope(run_subscription(connection, item))


@router.get("/literature")
def search_literature_library(
    q: str = Query(min_length=1, max_length=500),
    limit: int = Query(default=10, ge=1, le=50),
    accepted_only: bool = Query(default=False),
    connection: sqlite3.Connection = Depends(get_connection),
    _user: dict = Depends(get_current_user),
):
    items = literature.search_library(
        connection,
        q,
        limit=limit,
        accepted_only=accepted_only,
    )
    return envelope({"items": items, "total": len(items), "query": q})


@router.get("/literature/claims")
def list_literature_claims(
    review_status: str | None = Query(
        default=None,
        pattern="^(accepted|rejected|pending_review)$",
    ),
    limit: int = Query(default=50, ge=1, le=200),
    connection: sqlite3.Connection = Depends(get_connection),
    _user: dict = Depends(get_current_user),
):
    items = literature.list_claims(
        connection,
        review_status=review_status,
        limit=limit,
    )
    return envelope({"items": items, "total": len(items), "review_status": review_status})


@router.post("/literature/relations/detect")
def detect_literature_relations(
    payload: ClaimRelationDetectRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    _admin: dict = Depends(require_role("admin")),
):
    from ..services.literature_ingestion import detect_claim_relations_with_llm

    try:
        result = detect_claim_relations_with_llm(
            connection,
            limit=payload.limit,
            accepted_only=payload.accepted_only,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail=f"claim_relation_detection_failed:{str(exc)[:300]}",
        ) from exc
    return envelope(result)


@router.get("/literature/relations")
def list_literature_relations(
    review_status: str | None = Query(
        default=None,
        pattern="^(accepted|rejected|pending_review)$",
    ),
    limit: int = Query(default=100, ge=1, le=500),
    connection: sqlite3.Connection = Depends(get_connection),
    _user: dict = Depends(get_current_user),
):
    items = literature.list_relations(
        connection,
        review_status=review_status,
        limit=limit,
    )
    return envelope({"items": items, "total": len(items), "review_status": review_status})


@router.patch("/literature/relations/{relation_id}")
def review_literature_relation(
    relation_id: str,
    payload: ClaimReviewRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(require_role("admin", "researcher")),
):
    item = literature.review_relation(
        connection,
        relation_id,
        review_status=payload.review_status,
        reviewed_by=user["user_id"],
    )
    if item is None:
        raise HTTPException(status_code=404, detail="claim_relation_not_found")
    return envelope(item)


@router.patch("/literature/claims/{claim_id}")
def review_literature_claim(
    claim_id: str,
    payload: ClaimReviewRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(require_role("admin", "researcher")),
):
    item = literature.review_claim(
        connection,
        claim_id,
        review_status=payload.review_status,
        reviewed_by=user["user_id"],
    )
    if item is None:
        raise HTTPException(status_code=404, detail="claim_not_found")
    return envelope(item)


@router.get("/literature/{document_id}")
def get_literature_document(
    document_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    _user: dict = Depends(get_current_user),
):
    item = literature.get_document(connection, document_id)
    if item is None:
        raise HTTPException(status_code=404, detail="literature_document_not_found")
    return envelope(item)


@router.get("/cluster/drafts")
def list_cluster_drafts(
    project_id: str | None = Query(default=None),
    _user: dict = Depends(get_current_user),
):
    return envelope({"items": cluster.list_drafts(project_id=project_id)})


@router.post("/cluster/drafts")
def create_cluster_draft(
    payload: ClusterDraftRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, payload.project_id)
    try:
        item = cluster.create_draft(
            project_id=payload.project_id,
            created_by=user["user_id"],
            job_name=payload.job_name,
            command=payload.command,
            queue=payload.queue,
            gpu_count=payload.gpu_count,
            cpu_count=payload.cpu_count,
            setup_lines=payload.setup_lines,
            expected_outputs=payload.expected_outputs,
            rationale=payload.rationale,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return envelope(item)


@router.get("/cluster/drafts/{draft_id}")
def get_cluster_draft(
    draft_id: str,
    _user: dict = Depends(get_current_user),
):
    try:
        return envelope(cluster.refresh_draft(draft_id))
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@router.post("/cluster/drafts/{draft_id}/confirm")
def confirm_cluster_draft(
    draft_id: str,
    user: dict = Depends(get_current_user),
):
    if user.get("role") == "viewer":
        raise HTTPException(status_code=403, detail="forbidden")
    try:
        return envelope(cluster.submit_draft(draft_id, confirmed_by=user["user_id"]))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.get("/cluster/drafts/{draft_id}/download")
def download_cluster_output(
    draft_id: str,
    path: str = Query(min_length=1, max_length=500),
    _user: dict = Depends(get_current_user),
):
    try:
        filename, content = cluster.download_output(draft_id, path)
    except ValueError as exc:
        status = 404 if str(exc) == "output_not_found" else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc
    return Response(
        content=content,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/route-plan")
def route_plan(
    payload: RoutePlanRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, payload.project_id)
    nodes = catalog.list_workflow_nodes(connection, "run_pd1_round1") if payload.project_id else []
    return envelope({
        "mode": "rule_based_demo",
        "route": [
            "target_intake",
            "RFdiffusion backbone generation",
            "ProteinMPNN sequence design",
            "AlphaFold2 complex prediction",
            "Rosetta relax / interface scoring",
            "BDA filters",
            "Wet-lab validation",
        ],
        "compute_status": "not_connected",
        "note": "Demo mode returns a precomputed PD-1 binder route.",
        "input_summary": payload.model_dump(),
        "existing_node_count": len(nodes),
    })


@router.post("/research-briefs")
def create_research_brief(
    payload: ResearchBriefCreateRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, payload.project_id)
    from ..services.sweet_protein_planner import create_brief

    item = create_brief(
        connection,
        project_id=payload.project_id,
        title=payload.title,
        objective=payload.objective,
        product_context=payload.product_context,
        constraints=payload.constraints,
        source_material=payload.source_material,
        created_by=user["user_id"],
    )
    return envelope(item)


@router.get("/research-briefs")
def list_research_briefs(
    project_id: str = Query(min_length=1, max_length=160),
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, project_id)
    items = research_planner.list_project_briefs(connection, project_id)
    return envelope({"items": items, "total": len(items)})


@router.get("/research-briefs/{research_brief_id}")
def get_research_brief(
    research_brief_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    item = research_planner.get_brief(connection, research_brief_id)
    if item is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, item["project_id"])
    return envelope({
        **item,
        "findings": research_planner.list_findings(connection, research_brief_id),
        "questions": research_execution.list_questions(connection, research_brief_id),
        "research_runs": research_execution.list_runs(connection, research_brief_id),
        "hypotheses": research_execution.list_hypotheses(connection, research_brief_id),
    })


@router.post("/research-briefs/{research_brief_id}/sources/markdown")
def ingest_research_markdown(
    research_brief_id: str,
    payload: MarkdownResearchSourceRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    brief = research_planner.get_brief(connection, research_brief_id)
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    from ..services.research_source_ingestion import ingest_markdown_source

    try:
        item = ingest_markdown_source(
            connection,
            research_brief_id=research_brief_id,
            title=payload.title,
            content=payload.content,
            source_uri=payload.source_uri,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return envelope(item)


@router.post("/research-briefs/{research_brief_id}/plan")
def create_research_plan(
    research_brief_id: str,
    payload: ResearchPlanRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    brief = research_planner.get_brief(connection, research_brief_id)
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    from ..services.sweet_protein_planner import generate_plan

    try:
        item = generate_plan(
            connection,
            research_brief_id=research_brief_id,
            selected_route=payload.selected_route,
            created_by=user["user_id"],
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return envelope(item)


@router.post("/research-briefs/{research_brief_id}/research-runs")
def create_research_run(
    research_brief_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    brief = research_planner.get_brief(connection, research_brief_id)
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    from ..services.research_execution_service import ensure_questions

    ensure_questions(connection, research_brief_id)
    run = research_execution.create_run(
        connection,
        research_run_id=f"research_run_{uuid.uuid4().hex[:12]}",
        research_brief_id=research_brief_id,
        created_by=user["user_id"],
    )
    return envelope(run)


@router.post("/research-briefs/{research_brief_id}/sequence-comparison")
def compare_research_sequences(
    research_brief_id: str,
    payload: SequenceComparisonRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    brief = research_planner.get_brief(connection, research_brief_id)
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    from ..services.protein_comparison_service import compare_sequences

    try:
        return envelope(compare_sequences(payload.sequences))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.post("/research-briefs/{research_brief_id}/structure-comparison")
def compare_research_structures(
    research_brief_id: str,
    payload: StructureComparisonRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    brief = research_planner.get_brief(connection, research_brief_id)
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    from ..services.protein_comparison_service import compare_structures

    try:
        return envelope(compare_structures(
            connection,
            project_id=brief["project_id"],
            artifact_ids=payload.artifact_ids,
        ))
    except ValueError as exc:
        detail = str(exc)
        status = 403 if detail == "artifact_project_mismatch" else 400
        raise HTTPException(status_code=status, detail=detail) from exc


@router.post("/research-runs/{research_run_id}/start")
def start_research_run(
    research_run_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    run = research_execution.get_run(connection, research_run_id)
    if run is None:
        raise HTTPException(status_code=404, detail="research_run_not_found")
    brief = research_planner.get_brief(connection, run["research_brief_id"])
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    if run["status"] in {"running", "completed", "partial"}:
        return envelope(run)
    from ..services.research_execution_service import execute_research_run

    return envelope(execute_research_run(connection, research_run_id))


@router.get("/research-runs/{research_run_id}")
def get_research_run(
    research_run_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    run = research_execution.get_run(connection, research_run_id)
    if run is None:
        raise HTTPException(status_code=404, detail="research_run_not_found")
    brief = research_planner.get_brief(connection, run["research_brief_id"])
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    return envelope(run)


@router.patch("/research-evidence/{evidence_link_id}")
def review_research_evidence(
    evidence_link_id: str,
    payload: EvidenceReviewRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    row = connection.execute(
        """
        SELECT e.*, b.project_id FROM evidence_links e
        JOIN research_runs r ON r.research_run_id = e.research_run_id
        JOIN research_briefs b ON b.research_brief_id = r.research_brief_id
        WHERE e.evidence_link_id = ?
        """,
        (evidence_link_id,),
    ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="evidence_link_not_found")
    _ensure_project_access(connection, user, row["project_id"])
    return envelope(research_execution.review_evidence(
        connection,
        evidence_link_id,
        payload.review_status,
        user["user_id"],
    ))


@router.patch("/research-findings/{research_finding_id}/review")
def review_research_finding(
    research_finding_id: str,
    payload: ResearchFindingReviewRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    row = connection.execute(
        """
        SELECT f.*, b.project_id FROM research_findings f
        JOIN research_briefs b ON b.research_brief_id = f.research_brief_id
        WHERE f.research_finding_id = ?
        """,
        (research_finding_id,),
    ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="research_finding_not_found")
    _ensure_project_access(connection, user, row["project_id"])
    connection.execute(
        """
        UPDATE research_findings
        SET review_status = ?, updated_at = CURRENT_TIMESTAMP
        WHERE research_finding_id = ?
        """,
        (payload.review_status, research_finding_id),
    )
    item = connection.execute(
        "SELECT * FROM research_findings WHERE research_finding_id = ?",
        (research_finding_id,),
    ).fetchone()
    from ..repositories.base import decode_row

    return envelope(decode_row(item))


@router.get("/research-briefs/{research_brief_id}/dossier-export")
def export_research_dossier(
    research_brief_id: str,
    format: str = Query(default="markdown", pattern="^(markdown|json)$"),
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    brief = research_planner.get_brief(connection, research_brief_id)
    if brief is None:
        raise HTTPException(status_code=404, detail="research_brief_not_found")
    _ensure_project_access(connection, user, brief["project_id"])
    runs = [
        research_execution.get_run(connection, item["research_run_id"])
        for item in research_execution.list_runs(connection, research_brief_id)
    ]
    payload = {
        "brief": brief,
        "questions": research_execution.list_questions(connection, research_brief_id),
        "findings": research_planner.list_findings(connection, research_brief_id),
        "hypotheses": research_execution.list_hypotheses(connection, research_brief_id),
        "runs": [item for item in runs if item is not None],
        "workflow_plans": research_planner.list_plans(connection, research_brief_id),
        "reproducibility": {
            "exported_at": datetime.now(timezone.utc).isoformat(),
            "includes": [
                "brief_constraints",
                "source_material",
                "questions",
                "findings",
                "evidence",
                "hypotheses",
                "workflow_plan_versions",
                "workflow_parameters",
                "decision_gates",
            ],
        },
    }
    if format == "json":
        return Response(
            content=json.dumps(payload, ensure_ascii=False, indent=2, default=str),
            media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{research_brief_id}.json"'},
        )
    lines = [
        f"# {brief['title']}",
        "",
        brief["objective"],
        "",
        "## Assumptions",
        *[f"- {item.get('key')}: {item.get('value')} ({item.get('status')})" for item in brief.get("assumptions_json", [])],
        "",
        "## Findings",
        *[
            f"- [{item.get('review_status')}] **{item['title']}**: {item['statement']}"
            for item in payload["findings"]
        ],
        "",
        "## Design hypotheses",
        *[
            f"- **{item['hypothesis']}** — {item.get('rationale') or ''}"
            for item in payload["hypotheses"]
        ],
        "",
        "## Workflow plan versions",
        *[
            f"- v{item.get('version', 1)} `{item['workflow_plan_id']}` · "
            f"{item.get('selected_route') or 'route not selected'} · {item.get('status')}"
            for item in payload["workflow_plans"]
        ],
    ]
    return Response(
        content="\n".join(lines),
        media_type="text/markdown",
        headers={"Content-Disposition": f'attachment; filename="{research_brief_id}.md"'},
    )


@router.get("/workflow-plans/{workflow_plan_id}")
def get_workflow_plan(
    workflow_plan_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    item = research_planner.get_plan(connection, workflow_plan_id)
    if item is None:
        raise HTTPException(status_code=404, detail="workflow_plan_not_found")
    _ensure_project_access(connection, user, item["project_id"])
    return envelope(item)


@router.get("/workflow-runs/{workflow_run_id}/experiment-plan")
def get_workflow_experiment_plan(
    workflow_run_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    project_id = catalog.get_workflow_run_project_id(connection, workflow_run_id)
    _ensure_project_access(connection, user, project_id)
    item = experiment_plans.get_by_workflow_run(connection, workflow_run_id)
    if item is None:
        raise HTTPException(status_code=404, detail="experiment_plan_not_found")
    return envelope(item)


@router.get("/workflow-runs/{workflow_run_id}/parameter-recommendations")
def get_workflow_parameter_recommendations(
    workflow_run_id: str,
    node_run_id: str | None = None,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    project_id = catalog.get_workflow_run_project_id(connection, workflow_run_id)
    _ensure_project_access(connection, user, project_id)
    node = catalog.get_workflow_node(connection, node_run_id) if node_run_id else None
    if node_run_id and (node is None or node.get("workflow_run_id") != workflow_run_id):
        raise HTTPException(status_code=404, detail="node_not_found")
    plan_row = connection.execute(
        """
        SELECT nodes_json FROM workflow_plans
        WHERE materialized_workflow_run_id = ?
        """,
        (workflow_run_id,),
    ).fetchone()
    plan_node_key = None
    if node and plan_row:
        planned_nodes = json.loads(plan_row["nodes_json"] or "[]")
        matched = next(
            (
                planned
                for planned in planned_nodes
                if planned.get("name") == node.get("node_name")
                and planned.get("model_name") == node.get("model_name")
            ),
            None,
        )
        plan_node_key = (matched or {}).get("key")
    rows = connection.execute(
        """
        SELECT r.* FROM parameter_recommendations r
        JOIN workflow_plans p ON p.workflow_plan_id = r.workflow_plan_id
        WHERE p.materialized_workflow_run_id = ?
          AND (? IS NULL OR r.node_key = ?)
        ORDER BY r.node_key, r.parameter_key
        """,
        (workflow_run_id, plan_node_key, plan_node_key),
    ).fetchall()
    from ..repositories.base import decode_rows

    current_parameters = (node or {}).get("parameters_json") or {}
    items = decode_rows(rows)
    for item in items:
        item["current_value"] = current_parameters.get(item["parameter_key"])
        item["differs_from_recommendation"] = (
            item["current_value"] != item.get("recommended_value_json")
        )
    return envelope({"items": items, "total": len(items)})


@router.get("/experiment-plans/{experiment_plan_id}/result-template")
def download_experiment_result_template(
    experiment_plan_id: str,
    format: str = Query(default="csv", pattern="^(csv|xlsx|json)$"),
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    item = experiment_plans.get_plan(connection, experiment_plan_id)
    if item is None:
        raise HTTPException(status_code=404, detail="experiment_plan_not_found")
    _ensure_project_access(connection, user, item["project_id"])
    columns = (item.get("result_template_json") or {}).get("required_columns") or [
        "candidate_id", "stage_key", "metric", "value", "unit", "pass_status", "notes",
    ]
    if format == "json":
        return Response(
            content=json.dumps(
                {
                    "experiment_plan_id": experiment_plan_id,
                    "columns": columns,
                    "stage_keys": [step["stage_key"] for step in item["steps"]],
                },
                ensure_ascii=False,
                indent=2,
            ),
            media_type="application/json",
            headers={"Content-Disposition": f'attachment; filename="{experiment_plan_id}-results.json"'},
        )
    if format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "results"
        sheet.append(columns)
        for step in item["steps"]:
            row = {column: "" for column in columns}
            row["stage_key"] = step["stage_key"]
            sheet.append([row[column] for column in columns])
        instructions = workbook.create_sheet("instructions")
        instructions.append(["experiment_plan_id", experiment_plan_id])
        instructions.append(["status", item["status"]])
        instructions.append([])
        instructions.append(["stage_key", "title", "safety_level", "dependencies"])
        for step in item["steps"]:
            instructions.append([
                step["stage_key"],
                step["title"],
                step["safety_level"],
                ", ".join(step["dependencies_json"]),
            ])
        stream = io.BytesIO()
        workbook.save(stream)
        return Response(
            content=stream.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{experiment_plan_id}-results.xlsx"'},
        )
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(columns)
    for step in item["steps"]:
        row = {column: "" for column in columns}
        row["stage_key"] = step["stage_key"]
        writer.writerow([row[column] for column in columns])
    return Response(
        content=stream.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{experiment_plan_id}-results.csv"'},
    )


@router.patch("/experiment-plans/{experiment_plan_id}")
def update_experiment_plan(
    experiment_plan_id: str,
    payload: ExperimentPlanUpdateRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    item = experiment_plans.get_plan(connection, experiment_plan_id)
    if item is None:
        raise HTTPException(status_code=404, detail="experiment_plan_not_found")
    _ensure_project_access(connection, user, item["project_id"])
    if payload.status == "completed" and any(
        step["status"] != "completed" for step in item["steps"]
    ):
        raise HTTPException(
            status_code=409,
            detail="experiment_plan_completion_requires_all_steps",
        )
    return envelope(experiment_plans.update_plan(
        connection,
        experiment_plan_id,
        **payload.model_dump(),
    ))


@router.patch("/experiment-plan-steps/{step_id}")
def update_experiment_plan_step(
    step_id: str,
    payload: ExperimentStepUpdateRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    row = connection.execute(
        """
        SELECT s.*, p.project_id FROM experiment_plan_steps s
        JOIN experiment_plans p ON p.experiment_plan_id = s.experiment_plan_id
        WHERE s.experiment_plan_step_id = ?
        """,
        (step_id,),
    ).fetchone()
    if row is None:
        raise HTTPException(status_code=404, detail="experiment_plan_step_not_found")
    _ensure_project_access(connection, user, row["project_id"])
    if payload.result_artifact_id:
        artifact = connection.execute(
            "SELECT project_id FROM artifacts WHERE artifact_id = ?",
            (payload.result_artifact_id,),
        ).fetchone()
        if artifact is None or artifact["project_id"] != row["project_id"]:
            raise HTTPException(status_code=400, detail="experiment_result_artifact_mismatch")
    requested_status = payload.status
    existing_result_artifact_id = row["result_artifact_id"]
    if (
        requested_status == "completed"
        and not payload.result_artifact_id
        and not existing_result_artifact_id
    ):
        raise HTTPException(
            status_code=409,
            detail="experiment_step_completion_requires_result_artifact",
        )
    dependencies = json.loads(row["dependencies_json"] or "[]")
    if requested_status in {"ready", "in_progress", "completed"} and dependencies:
        placeholders = ",".join("?" for _ in dependencies)
        dependency_rows = connection.execute(
            f"""
            SELECT stage_key, status FROM experiment_plan_steps
            WHERE experiment_plan_id = ? AND stage_key IN ({placeholders})
            """,
            (row["experiment_plan_id"], *dependencies),
        ).fetchall()
        dependency_status = {item["stage_key"]: item["status"] for item in dependency_rows}
        incomplete = [key for key in dependencies if dependency_status.get(key) != "completed"]
        if incomplete:
            raise HTTPException(
                status_code=409,
                detail=f"experiment_dependencies_incomplete:{','.join(incomplete)}",
            )
    updated_step = experiment_plans.update_step(
        connection,
        step_id,
        {key: value for key, value in payload.model_dump().items() if value is not None},
    )
    plan = experiment_plans.synchronize_plan_status(connection, row["experiment_plan_id"])
    if plan and plan["status"] == "completed" and plan.get("node_run_id"):
        catalog.update_workflow_node(connection, plan["node_run_id"], status="completed")
        from ..services.run_coordinator import evaluate_downstream_nodes

        evaluate_downstream_nodes(
            connection,
            workflow_run_id=plan["workflow_run_id"],
            completed_node_run_id=plan["node_run_id"],
        )
        try:
            from ..services.campaign_service import sync_round_status

            sync_round_status(connection, plan["workflow_run_id"])
        except ValueError:
            pass
    return envelope({"step": updated_step, "plan": plan})


@router.get("/notifications")
def list_user_notifications(
    project_id: str | None = None,
    unread_only: bool = False,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    if project_id:
        _ensure_project_access(connection, user, project_id)
    items = automation.list_notifications(
        connection,
        user_id=user["user_id"],
        project_id=project_id,
        unread_only=unread_only,
    )
    items = [
        item for item in items
        if not item.get("project_id")
        or verify_project_access(connection, user, item["project_id"])
    ]
    return envelope({"items": items, "total": len(items)})


@router.patch("/notifications/{notification_id}/read")
def read_notification(
    notification_id: str,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    existing = connection.execute(
        "SELECT * FROM notifications WHERE notification_id = ?",
        (notification_id,),
    ).fetchone()
    if existing is None:
        raise HTTPException(status_code=404, detail="notification_not_found")
    if existing["project_id"]:
        _ensure_project_access(connection, user, existing["project_id"])
    if existing["user_id"] != user["user_id"]:
        raise HTTPException(status_code=404, detail="notification_not_found")
    item = automation.mark_read(connection, notification_id, user["user_id"])
    if item is None:
        raise HTTPException(status_code=404, detail="notification_not_found")
    return envelope(item)


@router.post("/workflow-plans/{workflow_plan_id}/materialize")
def materialize_workflow_plan(
    workflow_plan_id: str,
    payload: WorkflowPlanMaterializeRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    item = research_planner.get_plan(connection, workflow_plan_id)
    if item is None:
        raise HTTPException(status_code=404, detail="workflow_plan_not_found")
    _ensure_project_access(connection, user, item["project_id"])
    from ..services.sweet_protein_planner import materialize_plan

    try:
        result = materialize_plan(
            connection,
            workflow_plan_id=workflow_plan_id,
            selected_route=payload.selected_route,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return envelope(result)


@router.post("/candidate-explanation")
def candidate_explanation(
    payload: CandidateExplanationRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    _user: dict = Depends(get_current_user),
):
    candidate = catalog.get_candidate(connection, payload.candidate_id)
    reasons = [
        "Strong interface score and pLDDT in the demo batch.",
        "Acceptable SEC profile compared with aggregation-prone families.",
    ]
    recommendation = f"Review {payload.candidate_id} against BLI/SEC evidence before ordering."
    if candidate and candidate.get("decision") == "Anchor":
        recommendation = (
            f"Use {payload.candidate_id} as the round-two motif anchor when BLI/SEC evidence supports it."
        )
        if candidate.get("pred_kd"):
            reasons.insert(0, f"Predicted binding strength: {candidate['pred_kd']}.")
    return envelope({
        "candidate_id": payload.candidate_id,
        "recommendation": recommendation,
        "reasons": reasons,
    })


@router.post("/chat")
def copilot_chat(
    payload: CopilotChatRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, payload.project_id)
    return envelope(_resolve_chat(connection, payload))


@router.post("/chat/stream")
async def copilot_chat_stream(
    payload: CopilotChatRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, payload.project_id)
    result = await run_in_threadpool(_resolve_chat, connection, payload)

    async def event_generator():
        message = result.get("message", "")
        chunk_size = 24
        for i in range(0, len(message), chunk_size):
            yield {"event": "message", "data": message[i : i + chunk_size]}
        yield {"event": "done", "data": json.dumps({"skill_used": result.get("skill_used"), "mode": result.get("mode")})}

    return EventSourceResponse(event_generator())


@router.post("/chat/sync")
def copilot_chat_sync(
    payload: CopilotChatRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, payload.project_id)
    return envelope(_resolve_chat(connection, payload))


@router.post("/result-interpretation")
def result_interpretation(
    payload: ResultInterpretationRequest,
    connection: sqlite3.Connection = Depends(get_connection),
    user: dict = Depends(get_current_user),
):
    _ensure_project_access(connection, user, payload.project_id)
    summary = catalog.get_project_results_summary(connection, payload.project_id)
    package = catalog.get_project_delivery_package(connection, payload.project_id)
    constraints = (package or {}).get("redesign_constraints") or {}
    if isinstance(constraints, str):
        constraints = json.loads(constraints)
    return envelope({
        "project_id": payload.project_id,
        "summary": f"{summary['hit_rate_label']}; best BLI Kd is {summary['best_kd']}.",
        "round_two_constraints": constraints or {
            "preserve_candidate": summary.get("best_kd_candidate"),
            "increase_scaffold_diversity": True,
            "penalize_exposed_hydrophobic_area": True,
        },
    })


@router.get("/skills")
def list_skills(_user: dict = Depends(get_current_user)):
    return envelope([
        PROGRAMMABLE_BIOMATERIALS_SKILL,
        {"name": "workflow-adjust", "description": "Adjust workflow parameters and constraints"},
        {"name": "result-interpret", "description": "Interpret BLI/SEC experiment results"},
        {"name": "query-candidates", "description": "Query and rank candidates"},
        {"name": "structure-explain", "description": "Explain structure and interface"},
        {"name": "paper-reader", "description": "Literature search and summarization"},
    ])
